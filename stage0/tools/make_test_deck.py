#!/usr/bin/env python3
"""Generate a small, reproducible classroom PPTX. Development tool, not an iPad dependency."""
from io import BytesIO
from pathlib import Path
import argparse

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt


def make_deck(output: Path):
    if output.exists():
        raise ValueError('请使用新文件名，避免覆盖已有PPTX。')
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    entries = [
        ('空备注测试页', 'What color is it?\nA. Green   B. Red', ''),
        ('长备注测试页', 'Tell a short story about the forest.', '教师测试长备注。' * 400),
        ('图文测试页', 'Welcome to the forest!\nLook, think and answer.',
         'PRIVATE-FIXTURE-8642：仅供教师看的合成测试台词。')
    ]
    for title, body, notes in entries:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_frame = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11), Inches(1)).text_frame
        title_frame.text = title
        title_frame.paragraphs[0].runs[0].font.size = Pt(32)
        body_frame = slide.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(11), Inches(1.5)).text_frame
        body_frame.text = body
        for paragraph in body_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(28)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        if title == '图文测试页':
            image = Image.new('RGB', (640, 320), '#0d6259')
            drawing = ImageDraw.Draw(image)
            drawing.rectangle((30, 30, 610, 290), outline='white', width=8)
            drawing.text((270, 145), 'SAMPLE', fill='white')
            buffer = BytesIO()
            image.save(buffer, format='PNG')
            buffer.seek(0)
            slide.shapes.add_picture(buffer, Inches(1), Inches(1.7), width=Inches(6))
    # Keep package filenames different from display order to exercise true ordering.
    order = presentation.slides._sldIdLst
    order.insert(0, order[-1])
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    print(f'已生成3页测试绘本：{output}（{output.stat().st_size}字节）')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    make_deck(args.output)
